from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Colors
DARK_BLUE = RGBColor(10, 30, 80)
ORANGE = RGBColor(255, 120, 0)
GRAY = RGBColor(128, 128, 128)
WHITE = RGBColor(255, 255, 255)

def add_slide_with_content(prs, title_text, content_paragraphs, title_color):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.clear()
    
    for i, p_data in enumerate(content_paragraphs):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = p_data['text']
        p.level = p_data.get('level', 0)
        p.font.size = Pt(14) if p.level > 0 else Pt(18)
        p.font.bold = p_data.get('bold', False)

    return slide

# 1. Title Slide
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "BuildIQ\nConstruction Material Management System"
title.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
title.text_frame.paragraphs[0].font.bold = True

subtitle.text = ("A digital platform for efficient construction inventory tracking and reporting.\n\n"
                 "Team Members: [Names]\n"
                 "College Name: [College]\n"
                 "Guide Name: [Guide]")
subtitle.text_frame.paragraphs[0].font.color.rgb = GRAY

# 2. Introduction
intro_content = [
    {"text": "The construction industry is one of the largest sectors globally, but it often struggles with efficient resource management. Materials like cement, steel, and sand are the backbone of any project.", "level": 0},
    {"text": "Why Material Management is Crucial:", "level": 0, "bold": True},
    {"text": "Without proper management, projects experience significant delays and cost overruns. Keeping track of resources ensures smooth operations.", "level": 1},
    {"text": "Problems with Manual Management:", "level": 0, "bold": True},
    {"text": "Relying on paper ledgers or simple spreadsheets leads to human errors, lost data, and a lack of real-time visibility.", "level": 1},
    {"text": "The Need for Digital Solutions:", "level": 0, "bold": True},
    {"text": "A centralized digital platform automates tracking, minimizes mistakes, and allows project managers to monitor their sites efficiently from anywhere, bringing much-needed transparency to the entire construction process.", "level": 1}
]
add_slide_with_content(prs, "Introduction", intro_content, ORANGE)

# 3. Problem Statement
prob_content = [
    {"text": "Construction projects face numerous challenges when material tracking is done manually.", "level": 0},
    {"text": "Stock Shortages and Excess Inventory:", "level": 0, "bold": True},
    {"text": "Without real-time data, sites often run out of crucial materials, pausing work. Conversely, over-ordering ties up capital and leads to material degradation.", "level": 1},
    {"text": "Material Wastage and Theft:", "level": 0, "bold": True},
    {"text": "Poor tracking allows materials to be wasted or stolen without anyone noticing until it's too late.", "level": 1},
    {"text": "Difficulty Tracking Across Branches:", "level": 0, "bold": True},
    {"text": "Large companies struggle to move resources between multiple sites efficiently, leading to miscommunication and delayed deliveries.", "level": 1},
    {"text": "Reporting and Decision-Making Issues:", "level": 0, "bold": True},
    {"text": "Generating daily or weekly reports from manual records takes hours. Managers lack the accurate, timely insights required to make critical financial and operational decisions.", "level": 1}
]
add_slide_with_content(prs, "Problem Statement", prob_content, DARK_BLUE)

# 4. Purpose of the Project
purp_content = [
    {"text": "The primary idea behind BuildIQ is to revolutionize how construction companies handle their resources.", "level": 0},
    {"text": "Why It Was Developed:", "level": 0, "bold": True},
    {"text": "BuildIQ was built to eliminate the chaos of paper-based tracking. It replaces outdated methods with a fast, reliable, and user-friendly web application.", "level": 1},
    {"text": "Centralized Stock Management:", "level": 0, "bold": True},
    {"text": "The system acts as a single source of truth. All materials stored in the central warehouse are logged digitally, providing absolute clarity on what is available at any moment.", "level": 1},
    {"text": "Improving Communication:", "level": 0, "bold": True},
    {"text": "It bridges the gap between the main warehouse and branch sites. Constructors can instantly request materials through the platform, and managers can approve and dispatch them without endless phone calls or paperwork.", "level": 1}
]
add_slide_with_content(prs, "Purpose of the Project", purp_content, ORANGE)

# 5. Objectives
obj_content = [
    {"text": "BuildIQ aims to achieve several critical goals to streamline operations:", "level": 0},
    {"text": "Efficient Inventory Management:", "level": 0, "bold": True},
    {"text": "Maintain an accurate, real-time database of all materials across all sites, ensuring nothing is lost.", "level": 1},
    {"text": "Consumption Tracking:", "level": 0, "bold": True},
    {"text": "Monitor exactly how much of each material is used daily, identifying patterns and preventing unexpected shortages.", "level": 1},
    {"text": "Report Generation:", "level": 0, "bold": True},
    {"text": "Automatically create professional PDF reports for audits, billing, and management reviews with a single click.", "level": 1},
    {"text": "Demand Forecasting:", "level": 0, "bold": True},
    {"text": "Use historical consumption data to predict future needs, allowing better budget planning.", "level": 1},
    {"text": "Transparency and Accountability:", "level": 0, "bold": True},
    {"text": "Assign clear roles and track every action, ensuring that every bag of cement and piece of steel is fully accounted for by a specific user.", "level": 1}
]
add_slide_with_content(prs, "Objectives", obj_content, DARK_BLUE)

# 6. How the Application Was Developed
dev_content = [
    {"text": "BuildIQ is built using the MERN stack, chosen for its speed, scalability, and robust community support.", "level": 0},
    {"text": "Frontend - React.js:", "level": 0, "bold": True},
    {"text": "React was used to build a dynamic, fast, and responsive user interface. It allows for seamless navigation and real-time dashboard updates without reloading the page.", "level": 1},
    {"text": "Backend - Node.js and Express.js:", "level": 0, "bold": True},
    {"text": "Node.js provides a powerful server environment, while Express.js creates efficient APIs. They handle all logic, from processing material requests to authenticating users.", "level": 1},
    {"text": "Database - MongoDB:", "level": 0, "bold": True},
    {"text": "As a NoSQL database, MongoDB offers flexibility in storing complex data like varying material types, user logs, and transaction histories efficiently.", "level": 1},
    {"text": "Working Together:", "level": 0, "bold": True},
    {"text": "The frontend sends requests to the Node backend, which securely accesses and updates the MongoDB database, returning real-time data back to the user's dashboard.", "level": 1}
]
add_slide_with_content(prs, "How the Application Was Developed", dev_content, ORANGE)

# 7. System Architecture
arch_content = [
    {"text": "The system is structured around different roles, ensuring data security and a clear flow of operations.", "level": 0},
    {"text": "Admin and Stock Manager:", "level": 0, "bold": True},
    {"text": "The Admin oversees the entire system, managing users and global settings. The Stock Manager adds new stock to the central database and approves branch requests.", "level": 1},
    {"text": "Constructor/Branch Manager:", "level": 0, "bold": True},
    {"text": "Constructors log into their specific branch, request required materials from the central warehouse, and record daily consumption on-site.", "level": 1},
    {"text": "Data Flow:", "level": 0, "bold": True},
    {"text": "A Constructor submits a material request via the UI. The request is stored in the Database. The Stock Manager sees the request on their Dashboard, approves it, and the system automatically deducts the stock from the warehouse and adds it to the branch's inventory. Analytics are instantly updated.", "level": 1}
]
add_slide_with_content(prs, "System Architecture", arch_content, DARK_BLUE)

# 8. Key Features
feat_content = [
    {"text": "BuildIQ provides a comprehensive suite of tools for complete material control.", "level": 0},
    {"text": "Inventory Management & Material Requests:", "level": 0, "bold": True},
    {"text": "Users can categorize and search for items easily. Branch managers can submit standardized requests, preventing miscommunication over quantities.", "level": 1},
    {"text": "AI-Based Suggestions:", "level": 0, "bold": True},
    {"text": "When typing material names, the system smartly suggests standard items (e.g., 'Portland Cement 50kg'), reducing duplicate entries and keeping data clean.", "level": 1},
    {"text": "Analytics Dashboard & Reports:", "level": 0, "bold": True},
    {"text": "Interactive charts display consumption trends over time. Users can instantly generate detailed, professional PDF reports for any date range.", "level": 1},
    {"text": "Branch-Wise Tracking:", "level": 0, "bold": True},
    {"text": "Stock is not just tracked globally. Managers can view the exact inventory levels at Site A versus Site B, allowing for smart transfers between branches.", "level": 1}
]
add_slide_with_content(prs, "Key Features", feat_content, ORANGE)

# 9. How the System Works
works_content = [
    {"text": "Using BuildIQ is a seamless, step-by-step process designed for daily construction workflows.", "level": 0},
    {"text": "Step 1: Adding Stock:", "level": 0, "bold": True},
    {"text": "A delivery arrives at the main warehouse. The Stock Manager logs into the system, inputs the invoice details, and adds the materials to the central inventory.", "level": 1},
    {"text": "Step 2: Requesting Materials:", "level": 0, "bold": True},
    {"text": "A Constructor at Site A realizes they need 100 bags of cement for tomorrow. They open BuildIQ and submit a digital request.", "level": 1},
    {"text": "Step 3: Approval and Transfer:", "level": 0, "bold": True},
    {"text": "The Stock Manager approves the request. The system automatically updates the central database and credits Site A's digital inventory.", "level": 1},
    {"text": "Step 4: Consumption and Reporting:", "level": 0, "bold": True},
    {"text": "At the end of the day, the Constructor logs the materials used. The Dashboard updates its graphs, showing current stock levels and consumption rates instantly.", "level": 1}
]
add_slide_with_content(prs, "How the System Works", works_content, DARK_BLUE)

# 10. Benefits of the Application
ben_content = [
    {"text": "BuildIQ transforms traditional operations, delivering significant value to construction firms.", "level": 0},
    {"text": "Saves Time and Reduces Costs:", "level": 0, "bold": True},
    {"text": "By automating manual ledger entries and report generation, staff can focus on actual construction. Accurate tracking prevents over-purchasing and expensive emergency orders.", "level": 1},
    {"text": "Minimizes Wastage:", "level": 0, "bold": True},
    {"text": "Strict accountability and tracking ensure that materials are not lost, misplaced, or spoiled due to improper storage management.", "level": 1},
    {"text": "Improves Decision-Making:", "level": 0, "bold": True},
    {"text": "Project managers can view real-time data and charts. If they see steel consumption is higher than expected, they can investigate immediately rather than waiting for end-of-month audits.", "level": 1},
    {"text": "Increases Transparency:", "level": 0, "bold": True},
    {"text": "Every transaction is logged with a user ID and timestamp, eliminating disputes between the warehouse and branch sites.", "level": 1}
]
add_slide_with_content(prs, "Benefits of the Application", ben_content, ORANGE)

# 11. How Users Can Use the Application
use_content = [
    {"text": "The platform is designed with a user-friendly interface requiring minimal training.", "level": 0},
    {"text": "Secure Login Process:", "level": 0, "bold": True},
    {"text": "Users log in with their credentials. The system detects their role (Admin, Manager, Constructor) and loads the appropriate dashboard securely.", "level": 1},
    {"text": "Stock Management and Requests:", "level": 0, "bold": True},
    {"text": "Through an intuitive form, users can view a list of available materials. Requesting items takes just a few clicks—select the item, enter the quantity, and submit.", "level": 1},
    {"text": "Dashboard and Reports:", "level": 0, "bold": True},
    {"text": "Upon logging in, users are greeted with visual graphs of their recent activity. To generate a report, they navigate to the 'Reports' section, select a date range, and click 'Export PDF' to download.", "level": 1}
]
add_slide_with_content(prs, "How Users Can Use the Application", use_content, DARK_BLUE)

# 12. Future Enhancements
fut_content = [
    {"text": "BuildIQ is designed to scale and evolve. Several advanced features are planned for future versions.", "level": 0},
    {"text": "Mobile Application and QR Scanning:", "level": 0, "bold": True},
    {"text": "A dedicated mobile app will allow site workers to scan QR codes on material batches to log deliveries and consumption instantly from their phones.", "level": 1},
    {"text": "Predictive Analytics using AI:", "level": 0, "bold": True},
    {"text": "Machine learning algorithms will analyze past projects to automatically predict material requirements and alert managers before shortages occur.", "level": 1},
    {"text": "Real-Time Notifications and Cloud Deployment:", "level": 0, "bold": True},
    {"text": "Integration with SMS and email will provide instant alerts for low stock or approved requests. Full cloud deployment will allow multi-company, multi-tenant usage seamlessly.", "level": 1}
]
add_slide_with_content(prs, "Future Enhancements", fut_content, ORANGE)

# 13. Conclusion
conc_content = [
    {"text": "BuildIQ provides a necessary digital transformation for the construction sector.", "level": 0},
    {"text": "Summary of the Project:", "level": 0, "bold": True},
    {"text": "We have successfully developed a centralized, MERN-stack based platform that simplifies the complex task of managing construction materials across multiple sites.", "level": 1},
    {"text": "Impact on Construction Companies:", "level": 0, "bold": True},
    {"text": "By shifting from manual ledgers to a digital dashboard, companies can save countless hours, significantly reduce material wastage, and maintain strict control over their budgets.", "level": 1},
    {"text": "A Practical Solution:", "level": 0, "bold": True},
    {"text": "Ultimately, BuildIQ is not just a tracking tool; it is a comprehensive management system that brings efficiency, transparency, and data-driven decision-making to construction projects of any scale.", "level": 1}
]
add_slide_with_content(prs, "Conclusion", conc_content, DARK_BLUE)

# Save presentation
prs.save('BuildIQ_Presentation_Detailed.pptx')
print("Presentation successfully updated with extensive content: BuildIQ_Presentation_Detailed.pptx")
